"""Generate the fact sheet Faiaz writes his own chapters from.

Deliberately contains no interpretation. Every entry is a measurement, its
source file, and what the number literally counts. The reasoning about what
the numbers mean is what Rector's Directive 5/2019 Art. (5) requires the
author to supply, so this document does not supply it.

    python make_factsheet.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

import docx_builder as B
from results import Results

OUT = Path(__file__).resolve().parents[1] / "Fact sheet.docx"
REPO = Path(__file__).resolve().parents[1] / "freshkeeper"


def fact_table(document, rows, caption_text, widths=(4.6, 3.0, 6.9)):
    B.table(document, ["Fact", "Value", "Where it comes from"], rows,
            "Table", caption_text, widths=list(widths), font_size=9.0)


def main() -> int:
    R = Results()
    B.reset_registry()
    d = B.new_document()

    h = d.add_paragraph()
    run = h.add_run("Fact sheet for writing Chapters 2, 5 and 6 and the abstract")
    run.bold = True
    run.font.size = Pt(16)
    h.paragraph_format.space_after = Pt(10)

    B.para(d,
           "This sheet lists every measurement the thesis reports, what each "
           "number literally counts, and the file it is read from. There is no "
           "interpretation in it on purpose. Under Rector's Directive 5/2019 "
           "Art. (5), the results and the conclusions have to be formulated by "
           "you, so the reasoning is the part this sheet leaves empty.")
    B.para(d,
           "How to use it. Open this sheet and the figures in docs/figures, "
           "close the current draft, and write what you think the numbers "
           "show. Then check your sentences against the boundaries in the last "
           "section. Every value below is also in freshkeeper/results/, so you "
           "can open the raw file if you want to check one.", italic=True)

    # ---------------------------------------------------------------
    B.heading(d, "The dataset", 1)
    fact_table(d, [
        ["Total photographs", B.__dict__ and R.n(R.audit["corpus_size"]),
         "dataset_audit.json, corpus_size"],
        ["Fruit types", "8 (apple, banana, grape, guava, jujube, orange, "
         "pomegranate, strawberry)", "the dataset's own fruit_type column"],
        ["Label balance", "6,194 fresh, 6,141 rotten", "counted from the parquet files"],
        ["Licence", "CC-BY-4.0", "the dataset card on HuggingFace"],
        ["Training images", R.n(R.audit["split_counts"]["train"]),
         "dataset_audit.json, split_counts.train"],
        ["Validation images", R.n(R.audit["split_counts"]["val"]),
         "dataset_audit.json, split_counts.val"],
        ["Test images", R.n(R.audit["split_counts"]["test"]),
         "dataset_audit.json, split_counts.test"],
    ], "What the image dataset contains. All images and their fresh/rotten "
       "labels are real. Nothing here is generated.")

    B.heading(d, "The leakage audit", 2)
    fact_table(d, [
        ["Test images close to a training image, before the fix",
         R.pct(R.naive_leak), "dataset_audit.json, naive_split.test_leakage_fraction"],
        ["Median similarity to nearest training image, before the fix",
         f"{R.naive_median_nn:.3f}",
         "dataset_audit.json, naive_split.nearest_neighbour_percentiles.p50"],
        ["Similarity threshold used to call two images near duplicates",
         f"cosine {R.audit['similarity_threshold']}",
         "dataset_audit.json, similarity_threshold"],
        ["Pairs of images linked as near duplicates", R.n(R.audit["linking_pairs"]),
         "dataset_audit.json, linking_pairs"],
        ["Groups the 12,335 images collapse into", R.n(R.audit["n_groups"]),
         "dataset_audit.json, n_groups"],
        ["Largest group", f"{R.audit['largest_group']} images",
         "dataset_audit.json, largest_group"],
        ["Test images close to a training image, after the fix",
         R.pct(R.audit["residual_test_leakage_fraction"], 2),
         "dataset_audit.json, residual_test_leakage_fraction"],
        ["Byte-identical duplicate files found", "0",
         "every file hashed; none matched"],
    ], "The near duplicate audit. The 'before' figures come from the naive "
       "split; the 'after' figures from the group aware split.")

    B.heading(d, "The labelling anomaly", 2)
    fact_table(d, [
        ["Images labelled fresh apple whose filename says FreshOrange", "734",
         "counted from the dataset metadata"],
        ["Similarity of the fresh-apple group centre to rotten apple",
         f"{R.anomaly['fresh_apple_vs_rotten_apple']:.3f}",
         "dataset_audit.json, label_anomaly_centroid_cosine"],
        ["Similarity of the fresh-apple group centre to fresh orange",
         f"{R.anomaly['fresh_apple_vs_fresh_orange']:.3f}",
         "dataset_audit.json, label_anomaly_centroid_cosine"],
        ["Overlap between the fresh-apple and fresh-orange file sets", "none",
         "file hashes compared"],
    ], "The fresh-apple filename anomaly. The two similarity numbers are what "
       "settled which field to trust.")

    # ---------------------------------------------------------------
    B.heading(d, "The physical spoilage model", 1)
    rows = [[r["commodity"].capitalize(),
             f"{r['published_days']:.0f} d published, {r['modelled_days']:.1f} d modelled",
             f"error {r['error_percent']:+.1f}%"] for r in R.shelf["rows"]]
    B.table(d, ["Commodity", "Shelf life at 4 °C, 85% RH", "Difference"], rows,
            "Table", "Modelled against published shelf life. Source: "
            "shelf_life_validation.json, rows. The coefficients were fitted to "
            "the published values, so agreement here is not an independent "
            "check.", widths=[4.0, 6.5, 4.0], font_size=9.0)

    sw = R.shelf["temperature_sweep_days"]
    B.table(d, ["Commodity", "2 °C", "4 °C", "7 °C", "10 °C", "15 °C"],
            [[c.capitalize()] + [f"{sw[c][t]:.1f}" for t in ("2.0","4.0","7.0","10.0","15.0")]
             for c in sw],
            "Table",
            "Days to spoilage at different storage temperatures, in days. "
            "Source: shelf_life_validation.json, temperature_sweep_days. "
            "Temperature was never used in fitting, so this is the independent "
            "check.", widths=[4.0, 2.1, 2.1, 2.1, 2.1, 2.1], font_size=9.0)

    fact_table(d, [
        ["Largest error against published shelf lives",
         f"{max(abs(r['error_percent']) for r in R.shelf['rows']):.1f}%",
         "shelf_life_validation.json, computed over rows"],
        ["Marginal threshold", str(R.corpus["marginal_threshold"]),
         "sensor_corpus_meta.json, marginal_threshold"],
        ["Spoiled threshold", str(R.corpus["spoiled_threshold"]),
         "sensor_corpus_meta.json, spoiled_threshold"],
        ["Simulated fridge set points", "2.5, 4.0, 5.5 and 7.5 °C",
         "sensor_corpus_meta.json, fridge_population"],
        ["Simulated door openings per day", "6, 10, 16 and 24",
         "sensor_corpus_meta.json, fridge_population"],
    ], "Model settings. The two thresholds define what counts as marginal and "
       "spoiled everywhere in the thesis.")

    # ---------------------------------------------------------------
    B.heading(d, "Results, visual classifier", 1)
    B.para(d, "This is the only model in the thesis measured entirely on real "
              "data. Everything in section 4 depends on simulated sensor values.")
    fact_table(d, [
        ["Accuracy on the test split", R.f4(R.visual["accuracy"]),
         "evaluation.json, visual_cnn_binary.accuracy"],
        ["ROC AUC", R.f4(R.visual["auc"]), "evaluation.json, visual_cnn_binary.auc"],
        ["Macro F1", R.f4(R.visual["macro_f1"]), "evaluation.json, visual_cnn_binary.macro_f1"],
        ["Recall on the rotten class", R.f4(R.visual["rotten_recall"]),
         "evaluation.json, visual_cnn_binary.rotten_recall"],
        ["Precision on the rotten class", R.f4(R.visual["rotten_precision"]),
         "evaluation.json, visual_cnn_binary.rotten_precision"],
        ["Test images scored", R.n(R.visual["n"]), "evaluation.json, visual_cnn_binary.n"],
        ["Fresh items called rotten", str(R.visual["confusion_matrix"][0][1]),
         "evaluation.json, confusion_matrix"],
        ["Rotten items called fresh", str(R.visual["confusion_matrix"][1][0]),
         "evaluation.json, confusion_matrix"],
        ["Validation score on the naive (leaky) split",
         R.pct(R.naive_best_val) if R.naive_best_val else "n/a",
         "cnn_training_history_naive_split.json"],
        ["Validation score after the leakage fix", R.pct(R.stage2_best_val),
         "cnn_training_history.json"],
        ["Head training epochs", str(R.stage1_epochs), "cnn_training_history.json"],
        ["Fine tuning epochs", str(R.stage2_epochs), "cnn_training_history.json"],
    ], "Visual classifier, fresh versus rotten, on the leakage controlled test "
       "split.")

    B.table(d, ["Commodity", "Test images", "Accuracy", "AUC"],
            [[r["commodity"].capitalize(), str(r["n"]), R.f4(r["accuracy"]), R.f4(r["auc"])]
             for r in sorted(R.per_commodity, key=lambda x: -x["accuracy"])],
            "Table", "Accuracy per fruit. Source: per_commodity.json.",
            widths=[4.5, 3.5, 3.3, 3.2], font_size=9.0)

    # ---------------------------------------------------------------
    B.heading(d, "Results, three state classification", 1)
    B.para(d, "Every number in this section depends on simulated sensor "
              "values. See the boundaries in section 8 before writing about "
              "them.")
    B.table(d, ["Model", "Accuracy", "Macro F1", "Recall (spoiled)", "Precision (spoiled)"],
            [[label, R.f4(R.acc(m)), R.f4(R.macro_f1(m)),
              R.f4(R.spoiled_recall(m)), R.f4(R.spoiled_precision(m))]
             for label, m in (("Vision only", "vision_only"),
                              ("Sensor only", "sensor_only"),
                              ("Fusion, raw join", "fusion_raw"),
                              ("Fusion, projected", "fusion"))],
            "Table", "The four models on the test split. Source: "
            "evaluation.json.", widths=[4.4, 2.6, 2.6, 3.0, 3.0], font_size=9.0)

    fact_table(d, [
        ["Sensor only minus best fusion", f"{R.fusion_gap_points:.1f} points",
         "computed from evaluation.json"],
        ["Projected fusion minus raw fusion", f"{R.projection_gain_points:.1f} points",
         "computed from evaluation.json"],
        ["Vision only, marginal items classified correctly",
         f"{R.cm('vision_only')[1][1]} of {sum(R.cm('vision_only')[1])}",
         "evaluation.json, vision_only.confusion_matrix"],
        ["Vision only, marginal items called fresh", str(R.cm("vision_only")[1][0]),
         "evaluation.json, vision_only.confusion_matrix"],
        ["Vision only, marginal items called spoiled", str(R.cm("vision_only")[1][2]),
         "evaluation.json, vision_only.confusion_matrix"],
        ["Sensor only, spoiled items called fresh", str(R.cm("sensor_only")[2][0]),
         "evaluation.json, sensor_only.confusion_matrix"],
        ["Sensor only, fresh items called spoiled", str(R.cm("sensor_only")[0][2]),
         "evaluation.json, sensor_only.confusion_matrix"],
        ["Training samples", R.n(R.train["config"]["train_n"]),
         "fusion_training.json, config.train_n"],
        ["Validation samples", R.n(R.train["config"]["val_n"]),
         "fusion_training.json, config.val_n"],
    ], "Derived figures and error counts for the three state task.")

    rows = []
    for label, m in (("Sensor only", "sensor_only"), ("Fusion", "fusion")):
        for i, cls in enumerate(("fresh", "marginal", "spoiled")):
            rows.append([label if i == 0 else "", cls.capitalize(),
                         R.f4(R.per_class(m, cls, "precision")),
                         R.f4(R.per_class(m, cls, "recall")),
                         R.f4(R.per_class(m, cls, "f1")),
                         str(R.ev[m]["per_class"][cls]["support"])])
    B.table(d, ["Model", "Class", "Precision", "Recall", "F1", "Items"], rows,
            "Table", "Per class figures for the two strongest models. Source: "
            "evaluation.json, per_class.",
            widths=[3.4, 2.6, 2.6, 2.4, 2.2, 2.3], font_size=9.0)

    # ---------------------------------------------------------------
    B.heading(d, "Speed and size", 1)
    st = R.pipe["stages_ms"]
    B.table(d, ["Stage", "Mean (ms)", "Share of total"],
            [[n, f"{st[k]['mean']:.2f}", f"{100*st[k]['share_of_total']:.1f}%"]
             for n, k in (("JPEG decode and resize", "decode_resize"),
                          ("Preprocessing", "preprocess"),
                          ("MobileNetV2 backbone", "backbone"),
                          ("Fusion head", "fusion_head"),
                          ("Total per item", "total"))],
            "Table", "Inference cost per item on the development computer. "
            "Source: pipeline_benchmark.json, stages_ms.",
            widths=[6.0, 4.0, 4.5], font_size=9.0)

    fact_table(d, [
        ["Measured total per item", f"{R.stage('total'):.2f} ms",
         "pipeline_benchmark.json, stages_ms.total.mean"],
        ["Scaling factor used for the Raspberry Pi estimate", f"{R.pi_factor:.1f}",
         "Geekbench 6 single core medians, M1 vs Pi 4"],
        ["Estimated Raspberry Pi 4 time per item", f"{R.pi_ms:.0f} ms",
         "pipeline_benchmark.json, projected_pi4_total_ms"],
        ["Six slots as a share of a 30 minute cycle", R.pct(R.duty, 2),
         "pipeline_benchmark.json, cycle_budget_utilisation"],
        ["Benchmark runs", str(R.pipe["runs"]), "pipeline_benchmark.json, runs"],
    ], "Timing. The Raspberry Pi figure is an estimate. Nothing was timed on a "
       "Pi.")

    B.table(d, ["Version", "Size (MB)", "Test accuracy"],
            [["Keras", f"{R.keras_mb('fusion'):.3f}", R.f4(R.acc("fusion"))]] +
            [[f"TFLite {v}", f"{R.tfl('fusion', v, 'megabytes'):.3f}",
              R.f4(R.tfl("fusion", v, "test_accuracy"))]
             for v in ("float32", "float16", "int8")],
            "Table", "Model size after conversion. Source: "
            "tflite_benchmark.json.", widths=[5.0, 4.0, 5.0], font_size=9.0)

    # ---------------------------------------------------------------
    B.heading(d, "The software", 1)
    fact_table(d, [
        ["Lines of Python", R.n(R.loc), "test_metrics.json, python_lines"],
        ["Files tracked in git", str(R.files), "test_metrics.json, tracked_files"],
        ["Automated tests", str(R.tests), "test_metrics.json, tests_passed"],
        ["Test modules", str(R.test_modules), "test_metrics.json, test_modules"],
        ["Overall test coverage", f"{R.cov_total:.0f}%",
         "test_metrics.json, coverage_total_percent"],
        ["Coverage, alerts.py", f"{R.cov('freshkeeper/alerts.py')['percent']:.0f}%",
         "test_metrics.json, coverage_files"],
        ["Coverage, spoilage_model.py",
         f"{R.cov('freshkeeper/hardware/spoilage_model.py')['percent']:.0f}%",
         "test_metrics.json, coverage_files"],
        ["Coverage, rpi_backend.py",
         f"{R.cov('freshkeeper/hardware/rpi_backend.py')['percent']:.0f}%",
         "test_metrics.json, coverage_files"],
        ["Hardware cost", "EUR 147", "Table 2 in the thesis, retail prices"],
        ["Demo run length", f"{R.demo_days:.0f} simulated days in "
         f"{R.demo['cycles']} cycles", "demo_run.json"],
    ], "The implementation.")

    # ---------------------------------------------------------------
    B.heading(d, "Terms, in case you need to define them", 1)
    for term, meaning in [
        ("Spoilage extent", "a number from 0 to 1 for how far an item has gone "
         "towards being spoiled; 0 is perfect, 1 is fully spoiled"),
        ("Fresh, marginal, spoiled", "the three reported states; the cuts are at "
         "extent 0.35 and 0.70"),
        ("Freshness score", "the 0 to 100 number shown to the user, computed "
         "from the three class probabilities"),
        ("Macro F1", "the average of the F1 score of each class, counting each "
         "class equally no matter how common it is"),
        ("Recall on the spoiled class", "of all items that really are spoiled, "
         "the share the model calls spoiled"),
        ("Precision on the spoiled class", "of all items the model calls "
         "spoiled, the share that really are"),
        ("Near duplicate leakage", "when almost identical images end up in both "
         "the training and the test split, which makes the test score too high"),
        ("Ablation", "removing one input to see how much it was contributing"),
        ("Late fusion", "combining two inputs after each has gone through its "
         "own small network, rather than joining the raw inputs"),
    ]:
        B.rich_para(d, [(f"{term}. ", "b"), (meaning + ".", "")])

    # ---------------------------------------------------------------
    B.heading(d, "Boundaries: sentences you cannot write", 1)
    B.para(d, "These are not style advice. Each one is a claim the evidence "
              "does not support, and a reader who checks will find it.")
    B.reset_numbering()
    for claim in [
        "Anything of the form 'the sensors measured X'. No sensor was ever "
        "read. Every gas, temperature, humidity and mass value in the thesis "
        "comes from the physical model.",
        "Anything of the form 'the system reduced food waste by X'. No waste "
        "was measured. Published figures for other tools may be cited as "
        "other people's results, never as an outcome of this work.",
        "Anything of the form 'on a Raspberry Pi the system takes X ms'. The "
        "Pi figure is a host measurement multiplied by a benchmark ratio. Say "
        "estimated, and say where the ratio comes from.",
        "Anything of the form 'the interface is easy to use'. No usability "
        "study was run. You may say the design follows Nielsen's heuristics "
        "and that this was not tested.",
        "Anything of the form 'fusion works better than single sensors'. The "
        "measurement says the opposite under this simulation. If you argue it "
        "might hold on real hardware, mark it as a prediction.",
        "Anything of the form 'the model detects spoilage early'. The 98% is "
        "on a dataset with no in-between cases. It shows the easy case is "
        "easy.",
        "Anything of the form 'the drivers work'. They are written and their "
        "interface is tested. They have never been run on a Pi.",
    ]:
        B.numbered(d, claim)

    B.heading(d, "Where to find things", 1)
    B.table(d, ["What", "Where"],
            [["Raw result files", "freshkeeper/results/*.json"],
             ["Figures", "freshkeeper/docs/figures/"],
             ["The physics", "freshkeeper/freshkeeper/hardware/spoilage_model.py"],
             ["The dataset audit", "freshkeeper/scripts/audit_dataset.py"],
             ["Alert rules", "freshkeeper/freshkeeper/alerts.py"],
             ["Rerun everything", "cd freshkeeper && make all"],
             ["Rebuild the thesis", "cd build && python build.py"],
             ["Check numbers after editing", "cd build && python check_numbers.py \"../Main thesis.docx\""]],
            "Table", "Files and commands.", widths=[5.5, 9.0], font_size=9.0)

    d.save(str(OUT))
    print(f"Wrote {OUT}")
    print(f"  {len(d.paragraphs)} paragraphs, {len(d.tables)} tables")
    return 0


if __name__ == "__main__":
    sys.exit(main())
