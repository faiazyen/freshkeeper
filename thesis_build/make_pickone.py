"""A pick-one sheet for the technical conclusions the author must own.

For each question the author deferred on, this gives two or three honest
positions, all faithful to what the thesis actually measured. The author marks
the one they believe. Choosing is the author's own judgement, which is what
the results and conclusions require. The chosen lines then become the text,
with only grammar cleaned up.

    python make_pickone.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx.shared import Pt

import docx_builder as B
from results import Results

OUT = Path(__file__).resolve().parents[1] / "Pick one.docx"


def facts(document, text: str) -> None:
    p = document.add_paragraph()
    r = p.add_run("Facts: "); r.bold = True; r.font.size = Pt(9.5)
    r2 = p.add_run(text); r2.font.size = Pt(9.5); r2.italic = True
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = B.Cm(0.4)


def option(document, letter: str, text: str) -> None:
    p = document.add_paragraph()
    r = p.add_run(f"(  )  {letter})  "); r.bold = True
    p.add_run(text)
    p.paragraph_format.left_indent = B.Cm(0.4)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15


def own(document) -> None:
    p = document.add_paragraph()
    r = p.add_run("(  )  or write your own:"); r.bold = True
    p.paragraph_format.left_indent = B.Cm(0.4)
    p.paragraph_format.space_before = Pt(2)
    for _ in range(2):
        line = document.add_paragraph()
        run = line.add_run("_" * 88)
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        line.paragraph_format.left_indent = B.Cm(0.4)
        line.paragraph_format.space_after = Pt(6)


def main() -> int:
    R = Results()
    B.reset_registry()
    d = B.new_document()

    h = d.add_paragraph()
    run = h.add_run("Pick one: your conclusions, in your words")
    run.bold = True; run.font.size = Pt(16)
    h.paragraph_format.space_after = Pt(10)

    B.para(d,
           "Each question below has two or three answers. Every one is true and "
           "fair to what the thesis measured. Put an X in the box next to the "
           "one you actually believe. If none fits, write your own on the "
           "lines. This is quick, and the answer you choose is your own "
           "conclusion, which is what the university rules ask for.")
    B.para(d,
           "When you have marked them all, send it back. I will put your chosen "
           "answers into Chapters 5 and 6 and fix only the grammar.", italic=True)

    max_err = max(abs(r["error_percent"]) for r in R.shelf["rows"])

    # 1
    B.unnumbered_heading(d, "1. Is the physics model good enough to build on?", 1)
    facts(d, f"It matches published shelf lives within {max_err:.1f}%. But the "
             "numbers were fitted to those same shelf lives, so that match is "
             "not a real test. The temperature response, which was not fitted, "
             "did come out right.")
    option(d, "a", "Yes, good enough for a prototype. It matches published "
                   "shelf lives and reacts correctly to temperature, which was "
                   "not fitted. It is a fair base to build on, and testing "
                   "against real food is the next step.")
    option(d, "b", "Good enough to run the prototype, but not proven. Because "
                   "the numbers were fitted to known shelf lives, I cannot "
                   "claim it is correct until it is checked against real "
                   "measurements.")
    option(d, "c", "Only a starting point. More data and real testing are "
                   "needed before I would trust it.")
    own(d)

    # 2
    B.unnumbered_heading(d, "2. Does the 98% camera accuracy mean the system catches "
                 "spoilage early?", 1)
    facts(d, f"The camera scores {R.pct(R.visual['accuracy'])} on real photos. "
             "But the dataset only has clearly fresh and clearly rotten fruit. "
             "The hard case, food going bad that still looks fine, is not in "
             "the data.")
    option(d, "a", "No. The 98% only shows the camera can tell clearly fresh "
                   "from clearly rotten. It does not prove early detection, "
                   "because the in between case is not in the dataset.")
    option(d, "b", "It is a strong result on obvious cases, but not proof of "
                   "early detection. Catching spoilage early needs the sensor "
                   "data and real testing.")
    own(d)

    # 3
    B.unnumbered_heading(d, "3. Why did camera plus sensors (fusion) not beat sensors "
                 "alone? (the important one)", 1)
    facts(d, f"Sensors alone scored {R.pct(R.acc('sensor_only'))}. Fusion "
             f"scored {R.pct(R.acc('fusion'))}, which is lower. In this "
             "simulation the sensor numbers and the correct answers both come "
             "from the same model, so the sensors had an advantage.")
    option(d, "a", "In this test the sensor readings were made by the same "
                   "model that made the correct answers, so the sensors "
                   "already held the answer and the camera added little. On "
                   "real hardware, where sensors are noisier and less "
                   "complete, the camera could still help. This needs testing "
                   "on real hardware before deciding.")
    option(d, "b", "The camera and the sensors measured the same thing here, "
                   "so combining them was not worth the extra cost. In this "
                   "work, sensors alone were enough.")
    option(d, "c", "The result is not clear enough to decide. Fusion did not "
                   "help in this simulation, but the test was limited, so more "
                   "experiments are needed before I recommend for or against "
                   "it.")
    own(d)

    # 4
    B.unnumbered_heading(d, "4. Is a Raspberry Pi fast enough to run the system?", 1)
    facts(d, f"About {R.stage('total'):.0f} ms per item on the laptop. An "
             f"estimated {R.pi_ms:.0f} ms on a Pi, which is about "
             f"{R.pct(R.duty, 1)} of the thirty minute cycle. The Pi figure is "
             "an estimate, never measured on a real Pi.")
    option(d, "a", "Yes, easily. Even the estimated Pi time is a tiny part of "
                   "the cycle, so speed is not a problem. The one caution is "
                   "that this is an estimate, not measured on a real Pi.")
    option(d, "b", "Most likely yes, but I should not claim it until it is "
                   "measured on a real Pi. The estimate looks fine, with a "
                   "large margin.")
    own(d)

    # 5
    B.unnumbered_heading(d, "5. What is the lesson from the five bugs you found?", 1)
    facts(d, "Five bugs all produced believable but wrong output. None was "
             "found by reading the code. Each was found by checking the result "
             "against something outside the code, like a published number or a "
             "physical limit.")
    option(d, "a", "Believable output is not the same as correct output. The "
                   "bugs looked fine on screen and were only caught by "
                   "checking against outside facts. So every important result "
                   "needs an independent check, and each fix now has a test to "
                   "stop the bug coming back.")
    option(d, "b", "Testing matters more than careful coding. The bugs looked "
                   "fine until checked against real facts, so I added automatic "
                   "tests for each one.")
    own(d)

    # 6
    B.unnumbered_heading(d, "6. Is the system ready to be trusted, and was the work worth "
                 "doing?", 1)
    facts(d, "The sensors were simulated and never tested on hardware. But the "
             "whole thing can be checked and repeated, and it sets out the "
             "exact experiment that would prove or disprove it.")
    option(d, "a", "It is not ready to be trusted with real food yet, because "
                   "the sensors were simulated. But the work was worth doing, "
                   "because everything in it can be checked and repeated, and "
                   "it sets out exactly the experiment that would prove the "
                   "real system.")
    option(d, "b", "It is an early prototype, not ready for real use, but it "
                   "shows the idea can work and gives a clear plan for building "
                   "and testing the real thing.")
    own(d)

    B.unnumbered_heading(d, "When you are done", 1)
    B.para(d, "Send this back with your marks. I will drop your chosen answers "
              "into the thesis, fix only the grammar, rebuild, and run the "
              "checks. That closes the last real gap.")

    d.save(str(OUT))
    print(f"Wrote {OUT} ({len(d.paragraphs)} paragraphs)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
