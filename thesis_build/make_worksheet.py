"""A guided worksheet so the author writes Chapters 2, 5, 6 and the abstract
by answering short questions instead of facing a blank page.

Every fact is pre-filled. Every question asks for the author's own judgement
about what the facts mean. Stringing the answers together, in the author's own
rough English, becomes the chapter. Under Rector's Directive 5/2019 the
formulation of results and conclusions must be the author's (Art. 5); grammar
and style cleanup afterwards is allowed (Art. 4).

    python make_worksheet.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx.shared import Pt, RGBColor

import docx_builder as B
from results import Results

OUT = Path(__file__).resolve().parents[1] / "Worksheet.docx"
GREY = RGBColor(0x66, 0x66, 0x66)


def fact(document, text: str) -> None:
    p = document.add_paragraph()
    r = p.add_run("Facts you already have: ")
    r.bold = True; r.font.size = Pt(10)
    r2 = p.add_run(text); r2.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = B.Cm(0.4)


def question(document, text: str, lines: int = 3) -> None:
    p = document.add_paragraph()
    r = p.add_run("Q. "); r.bold = True
    p.add_run(text)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    for _ in range(lines):
        line = document.add_paragraph()
        run = line.add_run("_" * 92)
        run.font.color.rgb = GREY
        line.paragraph_format.space_after = Pt(6)
        line.paragraph_format.line_spacing = 1.0


def main() -> int:
    R = Results()
    B.reset_registry()
    d = B.new_document()

    h = d.add_paragraph()
    run = h.add_run("Worksheet: write your chapters by answering these questions")
    run.bold = True; run.font.size = Pt(16)
    h.paragraph_format.space_after = Pt(10)

    B.para(d,
           "This turns the four chapters you must write yourself into a set of "
           "short questions. Under each question the facts are already given, "
           "so you never have to remember a number. Answer each question in one "
           "to three sentences, in your own words. Rough English is fine and is "
           "in fact better, because it is yours.")
    B.para(d,
           "When you have answered every question, you will have written the "
           "chapters. Send the worksheet back and I will fix only grammar and "
           "spelling, without changing your meaning. That split is what the "
           "university rules allow: the thinking is yours, the language cleanup "
           "is help.", italic=True)
    B.para(d,
           "Do not copy the current draft. Close it. If you get stuck on a "
           "question, write what you would say out loud to a friend who asked "
           "you the same thing.", italic=True)

    # ---------------- Abstract ----------------
    B.heading(d, "Abstract (about 12 sentences)", 1)
    question(d, "In one or two sentences: why do people throw away food that is "
                "still edible? What is the problem this thesis is about?")
    question(d, "In one sentence: what did you build?")
    fact(d, f"visual classifier {R.pct(R.visual['accuracy'])}; physical model "
            f"within {max(abs(r['error_percent']) for r in R.shelf['rows']):.1f}% "
            f"of published shelf lives; sensor only {R.pct(R.acc('sensor_only'))} "
            f"vs fusion {R.pct(R.acc('fusion'))}.")
    question(d, "In your own words, what are the three main results, and what "
                "does each one tell you?", lines=5)
    question(d, "What is the one thing a reader must understand about the "
                "limits of this work? (Hint: the sensors were not real.)")

    # ---------------- Chapter 2 ----------------
    B.heading(d, "Chapter 2: Objectives and Methodology", 1)
    question(d, "Why did you choose this topic? What made it worth doing, for "
                "you personally?")
    B.para(d, "For each research question below, write in your own words why "
              "it matters. Do not just restate it.", italic=True)
    question(d, "RQ1 (can a physics model stand in for real spoilage data?): "
                "why does the answer matter for the rest of the thesis?", lines=2)
    question(d, "RQ2 (how well can a camera tell fresh from rotten, and what "
                "does that really prove?): why did you word the second half "
                "that way?", lines=2)
    question(d, "RQ3 (does combining camera and sensors beat either alone?): "
                "why was this worth testing even if the answer might be no?",
             lines=2)
    question(d, "RQ4 (is it fast enough for a Raspberry Pi?): why does this "
                "matter for a real product?", lines=2)
    question(d, "Objective 4 asked you to test whether fusion helps, knowing it "
                "might fail. Why were you willing to report a result that goes "
                "against your own design?")

    # ---------------- Chapter 5 ----------------
    B.heading(d, "Chapter 5: Results and Discussion", 1)

    B.heading(d, "The physical model", 2)
    fact(d, "The model matches published shelf lives within "
            f"{max(abs(r['error_percent']) for r in R.shelf['rows']):.1f}%. But "
            "the coefficients were fitted to those same shelf lives, so that "
            "match is not a real test. The real test is the temperature "
            "response, which was not fitted and still came out right.")
    question(d, "In your own words: does this convince you the model is good "
                "enough to build on? Why, or why not?")

    B.heading(d, "The visual classifier", 2)
    fact(d, f"{R.pct(R.visual['accuracy'])} accuracy, real photos. But the "
            "dataset only has clearly fresh and clearly rotten fruit. The "
            "in-between case is not in the data.")
    question(d, "What does this number actually tell you? Does it mean the "
                "system can catch spoilage early? Say what you think.", lines=4)

    B.heading(d, "Fusion versus sensor only", 2)
    fact(d, f"sensor only {R.pct(R.acc('sensor_only'))}, fusion "
            f"{R.pct(R.acc('fusion'))}. Fusion is worse. The reason: the "
            "simulated sensor numbers and the correct answers both come from "
            "the same model, so the sensors have an unfair advantage.")
    question(d, "In your own words, why do you think fusion lost? What does "
                "that mean for the design, and would you expect the same on "
                "real hardware?", lines=5)

    B.heading(d, "Speed", 2)
    fact(d, f"{R.stage('total'):.0f} ms per item on the laptop; an estimated "
            f"{R.pi_ms:.0f} ms on a Pi, which is about {R.pct(R.duty, 1)} of "
            "the 30 minute cycle.")
    question(d, "Is this fast enough for a Raspberry Pi? What is your "
                "judgement, and what is the catch with the estimate?", lines=3)

    B.heading(d, "The bugs you found", 2)
    fact(d, "Five bugs all produced believable but wrong output (the 2.3x "
            "shelf life error, the zeroed gas channel, training to chance, the "
            "mass delta, the changing random seed). Each was caught by "
            "comparing against something outside the code.")
    question(d, "What did you learn from the fact that none of these bugs was "
                "caught by reading the code? Why does that matter?", lines=3)

    # ---------------- Chapter 6 ----------------
    B.heading(d, "Chapter 6: Conclusion", 1)
    B.para(d, "Write your own one line answer to each research question. The "
              "facts are in Chapter 5; here just say yes or no and the reason.",
           italic=True)
    question(d, "RQ1, the physics model: your one line answer.", lines=2)
    question(d, "RQ2, the camera: your one line answer.", lines=2)
    question(d, "RQ3, fusion: your one line answer.", lines=2)
    question(d, "RQ4, speed: your one line answer.", lines=2)
    question(d, "What are you most proud of in this work? What is the real "
                "contribution, in your own words?", lines=4)
    question(d, "What is the biggest limitation, honestly?")
    question(d, "If you had six more months, what is the first thing you would "
                "do?")
    question(d, "Last sentence of the thesis: your honest take on whether this "
                "system is ready to be trusted, and why you still think the "
                "work was worth doing.", lines=3)

    B.heading(d, "When you are done", 1)
    B.para(d, "Save this file with your answers and send it back. I will turn "
              "your answers into finished paragraphs by fixing only grammar and "
              "spelling, and I will not change what you decided the results "
              "mean. Then we rebuild and run the checks.")

    d.save(str(OUT))
    print(f"Wrote {OUT} ({len(d.paragraphs)} paragraphs)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
