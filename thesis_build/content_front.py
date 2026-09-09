"""Front matter: title page, declaration, acknowledgement, abstract."""

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from docx_builder import add_toc, page_break, para
from results import Results

TITLE = "Design of an IoT-Based Food Spoilage Prediction System for Waste Reduction"
AUTHOR = "Faiaz Hossain Mazumder Yen"
SUPERVISOR = "Ing. Jiří Brožek, Ph.D."
DEPARTMENT = "Department of Information Engineering"
YEAR = "2026"

KEYWORDS = ("IoT, food waste reduction, spoilage prediction, sensor fusion, "
            "convolutional neural networks, predictive microbiology, Raspberry Pi, "
            "edge inference, Python, sustainable development")


def _centre(document, text, size, bold=False, space_after=12):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def build(document) -> None:
    R = Results()
    max_err = max(abs(r["error_percent"]) for r in R.shelf["rows"])

    _centre(document, "Czech University of Life Sciences Prague", 17, bold=True, space_after=6)
    _centre(document, "Faculty of Economics and Management", 14, space_after=4)
    _centre(document, DEPARTMENT, 13, space_after=90)
    _centre(document, "Bachelor Thesis", 20, bold=True, space_after=60)
    _centre(document, TITLE, 18, bold=True, space_after=50)
    _centre(document, AUTHOR, 14, space_after=6)
    _centre(document, f"Supervisor: {SUPERVISOR}", 12, space_after=90)
    _centre(document, f"© {YEAR} CZU Prague", 11, space_after=0)
    page_break(document)

    _centre(document, "Thesis Assignment", 14, bold=True, space_after=24)
    para(document,
         "Replace this page with the official thesis assignment exported to PDF "
         "from is.czu.cz (front and back page), as the faculty template requires.",
         italic=True, justify=False)
    page_break(document)

    _centre(document, "Declaration", 14, bold=True, space_after=18)
    para(document,
         f'I declare that I have worked on my bachelor thesis titled "{TITLE}" '
         "by myself and I have used only the sources mentioned at the end of the "
         "thesis. As the author of the bachelor thesis, I declare that the thesis "
         "does not break any copyrights.")
    para(document,
         "I declare that I have used AI tools in accordance with the "
         "university's internal regulations and principles of academic "
         "integrity and ethics.")
    para(document,
         "[Author to complete before submission. Rector's Directive 5/2019 "
         "Art. (6) requires a short statement here, in your own words, of how "
         "AI tools were used in this work as an auxiliary aid. Replace this "
         "note with that statement.]", italic=True)
    para(document,
         "The prototype software in this thesis was developed with AI help as an "
         "auxiliary tool for the research part. It is published under the MIT "
         "licence, and its full development history, including which parts were "
         "written with AI help, is public at https://github.com/faiazyen/freshkeeper. The image "
         "dataset used for training is a third party dataset under the CC-BY-4.0 "
         "licence. It is credited in Chapter 4 and in the references. The machine "
         "learning frameworks, libraries and the pretrained MobileNetV2 weights "
         "are third party components used under their open source licences.")
    para(document, "", space_after=36)
    para(document, "In Prague on ______________                "
                   "_________________________", justify=False)
    para(document, "                                                              "
                   f"      {AUTHOR}", justify=False)
    page_break(document)

    _centre(document, "Acknowledgement", 14, bold=True, space_after=18)
    para(document,
         f"I would like to thank {SUPERVISOR} for supervising this thesis, and "
         "especially for the review that led to this large revision. He said "
         "that the earlier draft described a system but did not show that the "
         "system existed. He was right. Acting on that comment changed the work "
         "a lot. A design document became a prototype that runs, with results "
         "that can be repeated and that are reported honestly, also when they "
         "are weaker than I hoped.")
    para(document,
         "I also thank the Department of Information Engineering for the "
         "teaching that made the practical part possible, and the people who "
         "maintain the open source projects this work is built on.")
    page_break(document)

    _centre(document, TITLE, 13, bold=True, space_after=18)
    _centre(document, "Abstract", 12, bold=True, space_after=10)
    para(document,
         "Households in rich countries throw away a large part of the food they "
         "buy. One common reason is that people are not sure if an item is still "
         "good, so they throw it away to be safe. Printed dates do not help "
         "much. A date describes an item stored under ideal conditions, not the "
         "item in a particular fridge. This thesis designs, builds and tests a "
         "prototype that measures the real storage conditions instead of "
         "guessing them.")
    para(document,
         "The system combines a camera, two gas sensors, a load cell and a "
         "temperature and humidity sensor on a Raspberry Pi 4. It classifies "
         "each monitored item as fresh, marginal or spoiled and shows the result "
         "in a web interface served from the device. All processing happens on "
         "the device. No image leaves the home network. The software is complete "
         f"and tested. It is {R.n(R.loc)} lines of Python and covers a physical "
         "spoilage model, a hardware layer with a real and a simulated backend, "
         "a machine learning pipeline, a REST API, a database and a web "
         f"interface. It has {R.tests} automated tests.")
    para(document,
         "Three results are reported. A MobileNetV2 classifier trained on "
         f"{R.n(R.audit['corpus_size'])} real photographs reaches "
         f"{R.pct(R.visual['accuracy'])} accuracy at telling fresh from rotten "
         "produce on a held out split. This came after an audit found and "
         "removed near duplicate leakage that had inflated an earlier number. A "
         "physical spoilage model built from the Ratkowsky and Gompertz "
         "equations reproduces published fridge shelf lives for eight fruits to "
         f"within {max_err:.1f}%. On the three state task, a sensor only "
         f"classifier reaches {R.pct(R.acc('sensor_only'))} and the fusion model "
         f"{R.pct(R.acc('fusion'))}. So fusion did not beat its own baseline. "
         "This negative result is explained, not hidden. The simulated sensor "
         "values and the true labels come from the same physical model, so the "
         "sensor branch has an unfair advantage. A fair test of fusion needs "
         "real hardware.")
    para(document,
         "The line between what was measured and what was simulated is stated "
         "everywhere a result depends on it. The images and their labels are "
         "real. The gas, temperature, humidity and mass readings are generated "
         "by the physical model. The Raspberry Pi drivers are written but not "
         "tested on hardware.")
    para(document, "", space_after=6)
    kp = document.add_paragraph()
    label = kp.add_run("Keywords: "); label.bold = True
    kp.add_run(KEYWORDS)
    kp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    page_break(document)

    ch = document.add_paragraph()
    run = ch.add_run("Table of Contents"); run.bold = True; run.font.size = Pt(16)
    ch.paragraph_format.space_after = Pt(14)
    add_toc(document)
    para(document,
         "To build this list in Word: click in it, press F9, and choose to "
         "update the whole table. In LibreOffice: Tools, Update, Indexes and "
         "Tables.", italic=True, space_after=4)
    page_break(document)
